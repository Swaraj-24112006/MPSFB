"""
Vendor Schedule Views for MPS Update Delivery Schedule tab.
All 9 views implementing:
1. Consolidated RM/PM Requirements Matrix (heavy computation read)
2. Single Schedule List & Create (CRUD)
3. Single Schedule Detail, Edit & Delete (CRUD with audit log)
4. Bulk Excel/CSV Schedule Upload
5. Auto-Fill Deficits (1-Click shortage coverage)
6. Change Log / History Drawer (Audit trail query)
7. Blank Excel Template Download
8. Pre-Filled Excel Template Download
9. Consolidated Matrix Excel Export
"""
import io
import csv
from datetime import datetime, date
from decimal import Decimal

from django.db import transaction
from django.db.models import Q
from django.http import StreamingHttpResponse
from django.utils import timezone
from rest_framework import status
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from rest_framework.views import APIView

from core.models import (
    VendorDeliverySchedule, DeliveryScheduleChangeLog, UploadBatch,
    WeekDefinition, BOMMaster, VendorSuppliedComponent, VendorBuyerMaster,
    RMPMComponentMaster
)
from core.serializers import (
    VendorDeliveryScheduleSerializer,
    DeliveryScheduleChangeLogSerializer
)
from core.services.consolidated_matrix_service import ConsolidatedMatrixService
from core.services.audit_log_service import AuditLogService
from core.services.vendor_schedule_week_resolver import VendorScheduleWeekResolver
from core.services.bulk_schedule_upload_service import (
    BulkScheduleUploadService, VALID_STATUSES
)
from core.services.auto_fill_deficit_service import AutoFillDeficitService


# ─────────────────────────────────────────────────────────────────────────────
# View 1: Consolidated RM/PM Requirements Matrix
# GET /api/vendor-schedule/consolidated-matrix/
# ─────────────────────────────────────────────────────────────────────────────

class ConsolidatedMatrixView(APIView):
    permission_classes = [AllowAny]

    def get(self, request):
        month = request.query_params.get('month')
        scope_mode = request.query_params.get('scope_mode', 'WEEK')
        week_code = request.query_params.get('week_code')
        buyer = request.query_params.get('buyer', 'ALL')
        category = request.query_params.get('category', 'ALL')
        criticality = request.query_params.get('criticality')
        search = request.query_params.get('search', '').strip()

        if not month:
            return Response({'error': 'month is required'}, status=status.HTTP_400_BAD_REQUEST)

        result = ConsolidatedMatrixService.compute(
            month=month,
            scope_mode=scope_mode,
            selected_week_code=week_code,
            buyer_filter=buyer,
            category_filter=category,
            criticality_filter=criticality,
            search_term=search or None,
        )
        return Response(result, status=status.HTTP_200_OK)


# ─────────────────────────────────────────────────────────────────────────────
# View 2: List and Create Vendor Delivery Schedules
# GET  /api/vendor-delivery-schedules/
# POST /api/vendor-delivery-schedules/
# ─────────────────────────────────────────────────────────────────────────────

class VendorScheduleListCreateView(APIView):
    permission_classes = [AllowAny]

    def get(self, request):
        month = request.query_params.get('month')
        week_code = request.query_params.get('week_code') or request.query_params.get('weekId')
        comp_code = request.query_params.get('component_code') or request.query_params.get('componentCode')
        vendor_code = request.query_params.get('vendor_code') or request.query_params.get('vendorCode')
        buyer_name = request.query_params.get('buyer_name') or request.query_params.get('buyerName')
        del_status = request.query_params.get('delivery_status') or request.query_params.get('deliveryStatus')
        search = request.query_params.get('search', '').strip()

        qs = VendorDeliverySchedule.objects.select_related('component', 'week').all().order_by('-expected_delivery_date')

        if month:
            week_codes = list(
                WeekDefinition.objects.filter(month=month).values_list('week_code', flat=True)
            )
            qs = qs.filter(week_id__in=week_codes)
        if week_code:
            qs = qs.filter(week_id=week_code)
        if comp_code:
            qs = qs.filter(component_id=comp_code)
        if vendor_code:
            qs = qs.filter(vendor_code=vendor_code)
        if buyer_name:
            qs = qs.filter(buyer_name=buyer_name)
        if del_status:
            qs = qs.filter(delivery_status=del_status)
        if search:
            qs = qs.filter(
                Q(po_number__icontains=search) |
                Q(component_id__icontains=search) |
                Q(vendor_name__icontains=search) |
                Q(buyer_name__icontains=search)
            )

        serializer = VendorDeliveryScheduleSerializer(qs, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)

    def post(self, request):
        """
        Create one vendor delivery schedule.
        Required: component_code, promised_qty, expected_delivery_date, changed_by, reason_for_change
        reason_for_change must be at least 10 characters.
        """
        data = request.data
        changed_by = data.get('changed_by', '').strip()
        reason = data.get('reason_for_change', '').strip()
        comp_code = (data.get('component_code') or data.get('componentCode', '')).strip()
        promised_qty = data.get('promised_qty') if 'promised_qty' in data else data.get('promisedQty', 0)
        delivery_date = data.get('expected_delivery_date') or data.get('expectedDeliveryDate')
        month = data.get('month', '')

        # Validation
        if not changed_by:
            return Response({'error': 'changed_by is required'}, status=status.HTTP_400_BAD_REQUEST)
        if len(reason) < 10:
            return Response(
                {'error': 'reason_for_change must be at least 10 characters'},
                status=status.HTTP_400_BAD_REQUEST
            )
        if not comp_code:
            return Response({'error': 'component_code is required'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            qty_val = float(promised_qty)
            if qty_val <= 0:
                raise ValueError
        except (ValueError, TypeError):
            return Response({'error': 'promised_qty must be greater than 0'}, status=status.HTTP_400_BAD_REQUEST)
        if not delivery_date:
            return Response({'error': 'expected_delivery_date is required'}, status=status.HTTP_400_BAD_REQUEST)

        # Parse delivery date
        try:
            parsed_date = datetime.strptime(str(delivery_date), '%Y-%m-%d').date()
        except ValueError:
            return Response(
                {'error': 'expected_delivery_date must be YYYY-MM-DD format'},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Resolve week_code from date
        if not month:
            month = str(parsed_date)[:7]
        week_code = VendorScheduleWeekResolver.resolve(parsed_date, month)

        # Validate delivery status enum
        delivery_status = data.get('delivery_status') or data.get('deliveryStatus', 'CONFIRMED_ON_TRACK')
        if delivery_status not in VALID_STATUSES:
            return Response(
                {'error': f'Invalid delivery_status. Must be one of: {sorted(VALID_STATUSES)}'},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Resolve vendor/buyer info if not provided
        vendor_code = data.get('vendor_code') or data.get('vendorCode', '')
        vendor_name = data.get('vendor_name') or data.get('vendorName', '')
        buyer_name = data.get('buyer_name') or data.get('buyerName', '')

        if not vendor_code or not vendor_name or not buyer_name:
            vsc = VendorSuppliedComponent.objects.filter(component_id=comp_code).select_related('vendor_buyer').first()
            if vsc and vsc.vendor_buyer:
                vb = vsc.vendor_buyer
                if not vendor_code:
                    vendor_code = vb.vendor_code
                if not vendor_name:
                    vendor_name = vb.vendor_name
                if not buyer_name:
                    buyer_name = vb.buyer_name

        po_number = (data.get('po_number') or data.get('poNumber', '')).strip()
        if not po_number:
            po_number = f"PO-{int(timezone.now().timestamp())}"

        # Fetch component description for richer audit logs
        comp_obj = RMPMComponentMaster.objects.filter(component_code=comp_code).first()
        comp_desc = comp_obj.component_description if comp_obj else ''

        from django.db import IntegrityError
        try:
            with transaction.atomic():
                sched = VendorDeliverySchedule.objects.create(
                    po_number=po_number,
                    component_id=comp_code,
                    vendor_code=vendor_code or 'V-NONE',
                    vendor_name=vendor_name,
                    buyer_name=buyer_name,
                    expected_delivery_date=parsed_date,
                    week_id=week_code,
                    promised_qty=Decimal(str(qty_val)),
                    carrier_or_tracking=data.get('carrier_or_tracking') or data.get('carrierOrTracking', ''),
                    delivery_status=delivery_status,
                    notes=data.get('notes', ''),
                )
                AuditLogService.log_create(sched, changed_by, reason, component_description=comp_desc)
        except IntegrityError:
            return Response(
                {'error': f'A delivery schedule for PO {po_number}, component {comp_code}, and date {parsed_date} already exists.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        serializer = VendorDeliveryScheduleSerializer(sched)
        return Response(serializer.data, status=status.HTTP_201_CREATED)


# ─────────────────────────────────────────────────────────────────────────────
# View 3: Retrieve, Update, Delete single schedule
# GET    /api/vendor-delivery-schedules/{id}/
# PATCH  /api/vendor-delivery-schedules/{id}/
# DELETE /api/vendor-delivery-schedules/{id}/
# ─────────────────────────────────────────────────────────────────────────────

class VendorScheduleDetailView(APIView):
    permission_classes = [AllowAny]

    def get_object(self, pk):
        try:
            return VendorDeliverySchedule.objects.select_related('component', 'week').get(pk=pk)
        except VendorDeliverySchedule.DoesNotExist:
            return None

    def get(self, request, pk):
        sched = self.get_object(pk)
        if not sched:
            return Response({'error': 'Schedule not found'}, status=status.HTTP_404_NOT_FOUND)
        return Response(VendorDeliveryScheduleSerializer(sched).data, status=status.HTTP_200_OK)

    def patch(self, request, pk):
        """
        Edit an existing delivery schedule.
        REQUIRED in payload: changed_by (str) + reason_for_change (str, min 10 chars)
        ALL other fields are optional.
        """
        sched = self.get_object(pk)
        if not sched:
            return Response({'error': 'Schedule not found'}, status=status.HTTP_404_NOT_FOUND)

        data = request.data
        changed_by = data.get('changed_by', '').strip()
        reason = data.get('reason_for_change', '').strip()

        if not changed_by:
            return Response({'error': 'changed_by is required'}, status=status.HTTP_400_BAD_REQUEST)
        if len(reason) < 10:
            return Response(
                {'error': 'reason_for_change must be at least 10 characters'},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Snapshot old values for audit log
        old_snapshot = {
            'promised_qty': float(sched.promised_qty),
            'expected_delivery_date': sched.expected_delivery_date.isoformat(),
            'delivery_status': sched.delivery_status,
            'vendor_name': sched.vendor_name,
            'buyer_name': sched.buyer_name,
            'carrier_or_tracking': sched.carrier_or_tracking or '',
            'notes': sched.notes or '',
        }

        # Apply updates
        field_map = {
            'po_number': ['po_number', 'poNumber'],
            'vendor_code': ['vendor_code', 'vendorCode'],
            'vendor_name': ['vendor_name', 'vendorName'],
            'buyer_name': ['buyer_name', 'buyerName'],
            'carrier_or_tracking': ['carrier_or_tracking', 'carrierOrTracking'],
            'delivery_status': ['delivery_status', 'deliveryStatus'],
            'notes': ['notes'],
        }

        for model_field, keys in field_map.items():
            for k in keys:
                if k in data:
                    setattr(sched, model_field, data[k])
                    break

        if 'component_code' in data or 'componentCode' in data:
            c_code = data.get('component_code') or data.get('componentCode')
            sched.component_id = c_code

        if 'promised_qty' in data or 'promisedQty' in data:
            raw_qty = data.get('promised_qty') if 'promised_qty' in data else data.get('promisedQty')
            try:
                val = float(raw_qty)
                if val <= 0:
                    raise ValueError
                sched.promised_qty = Decimal(str(val))
            except (ValueError, TypeError):
                return Response({'error': 'promised_qty must be greater than 0'}, status=status.HTTP_400_BAD_REQUEST)

        date_val = data.get('expected_delivery_date') or data.get('expectedDeliveryDate')
        if date_val:
            try:
                parsed_date = datetime.strptime(str(date_val), '%Y-%m-%d').date()
            except ValueError:
                return Response(
                    {'error': 'expected_delivery_date must be YYYY-MM-DD'},
                    status=status.HTTP_400_BAD_REQUEST
                )
            sched.expected_delivery_date = parsed_date
            month = str(parsed_date)[:7]
            sched.week_id = VendorScheduleWeekResolver.resolve(parsed_date, month)

        comp_obj = RMPMComponentMaster.objects.filter(component_code=sched.component_id).first()
        comp_desc = comp_obj.component_description if comp_obj else ''

        with transaction.atomic():
            sched.save()
            AuditLogService.log_update(sched, old_snapshot, changed_by, reason, component_description=comp_desc)

        return Response(VendorDeliveryScheduleSerializer(sched).data, status=status.HTTP_200_OK)

    def delete(self, request, pk):
        """
        Deletes a delivery schedule.
        REQUIRED in request body: changed_by + reason_for_change (min 10 chars)
        Writes audit log BEFORE deleting (FK goes SET NULL after delete).
        """
        sched = self.get_object(pk)
        if not sched:
            return Response({'error': 'Schedule not found'}, status=status.HTTP_404_NOT_FOUND)

        data = request.data
        changed_by = data.get('changed_by', '').strip()
        reason = data.get('reason_for_change', '').strip()

        if not changed_by:
            return Response({'error': 'changed_by is required'}, status=status.HTTP_400_BAD_REQUEST)
        if len(reason) < 10:
            return Response(
                {'error': 'reason_for_change must be at least 10 characters'},
                status=status.HTTP_400_BAD_REQUEST
            )

        comp_obj = RMPMComponentMaster.objects.filter(component_code=sched.component_id).first()
        comp_desc = comp_obj.component_description if comp_obj else ''

        with transaction.atomic():
            AuditLogService.log_delete(sched, changed_by, reason, component_description=comp_desc)
            sched.delete()

        return Response({'deleted': True}, status=status.HTTP_200_OK)


# ─────────────────────────────────────────────────────────────────────────────
# View 4: Bulk Excel/CSV Upload
# POST /api/uploads/vendor-delivery-schedules/
# ─────────────────────────────────────────────────────────────────────────────

class VendorScheduleBulkUploadView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        month = request.data.get('month')
        changed_by = request.data.get('changed_by', '').strip()
        import_mode = request.data.get('import_mode', 'APPEND')
        upload_scope = request.data.get('upload_scope', 'SELECTED_WEEK')
        week_code = request.data.get('week_code', '')
        pasted_text = request.data.get('pasted_text', '')
        file_obj = request.FILES.get('file')

        if not month:
            return Response({'error': 'month is required'}, status=status.HTTP_400_BAD_REQUEST)
        if not changed_by:
            return Response({'error': 'changed_by is required'}, status=status.HTTP_400_BAD_REQUEST)
        if not file_obj and not pasted_text.strip():
            return Response(
                {'error': 'Provide either a file or pasted_text'},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Verify weeks exist for month
        week_qs = WeekDefinition.objects.filter(month=month)
        if not week_qs.exists():
            return Response(
                {'error': f'No week definitions found for {month}. Define weeks first.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        first_week_code = week_qs.order_by('week_no').first().week_code

        # Pre-load BOM component codes and vendor-buyer mappings
        bom_component_codes = set(
            BOMMaster.objects.filter(is_active=True).values_list('component_id', flat=True)
        )
        vsc_rows = VendorSuppliedComponent.objects.select_related('vendor_buyer').all()
        vendor_by_component = {
            vsc.component_id: vsc.vendor_buyer for vsc in vsc_rows
        }

        # Parse file or pasted text
        file_name = 'Pasted Clipboard Spreadsheet Text'
        try:
            if file_obj:
                file_name = file_obj.name
                ext = file_name.rsplit('.', 1)[-1].lower()
                if ext in ('xlsx', 'xls'):
                    raw_rows = BulkScheduleUploadService.parse_file(file_obj)
                else:
                    text = file_obj.read().decode('utf-8', errors='replace')
                    raw_rows = BulkScheduleUploadService.parse_csv_text(text)
            else:
                raw_rows = BulkScheduleUploadService.parse_csv_text(pasted_text)
        except Exception as e:
            return Response({'error': f'Failed to parse file: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

        if not raw_rows:
            return Response({'error': 'No rows found in the uploaded file'}, status=status.HTTP_400_BAD_REQUEST)

        # Normalize each row
        parsed_rows = [
            BulkScheduleUploadService.normalize_row(
                raw, month, bom_component_codes,
                vendor_by_component, first_week_code
            )
            for raw in raw_rows
        ]

        # Create upload batch
        batch = UploadBatch.objects.create(
            upload_type='VENDOR_DELIVERY_SCHEDULE',
            uploaded_by=changed_by,
            uploaded_at=timezone.now(),
            file_name=file_name,
            minio_path='',
            status='PROCESSING',
            total_rows=len(raw_rows),
            imported_rows=0,
            error_rows=0,
        )

        result = BulkScheduleUploadService.apply(
            month=month,
            parsed_rows=parsed_rows,
            import_mode=import_mode,
            upload_scope=upload_scope,
            selected_week_code=week_code,
            changed_by=changed_by,
            upload_file_name=file_name,
            batch=batch,
        )

        return Response({
            'batch_id': batch.id,
            'month': month,
            'total_rows': len(raw_rows),
            'imported_rows': result['imported'],
            'error_rows': result['skipped'],
            'errors': result['errors'],
            'message': (
                f"Successfully imported {result['imported']} delivery schedule(s) "
                f"for {month}."
            ),
        }, status=status.HTTP_200_OK)


# ─────────────────────────────────────────────────────────────────────────────
# View 5: Auto-Fill All Deficits (1-Click)
# POST /api/vendor-delivery-schedules/auto-fill-deficits/
# ─────────────────────────────────────────────────────────────────────────────

class AutoFillDeficitsView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        month = request.data.get('month')
        scope_mode = request.data.get('scope_mode', 'WEEK')
        week_code = request.data.get('week_code')
        changed_by = request.data.get('changed_by', '').strip()

        if not month:
            return Response({'error': 'month is required'}, status=status.HTTP_400_BAD_REQUEST)
        if not changed_by:
            return Response({'error': 'changed_by is required'}, status=status.HTTP_400_BAD_REQUEST)

        result = AutoFillDeficitService.auto_fill(
            month=month,
            scope_mode=scope_mode,
            selected_week_code=week_code,
            changed_by=changed_by,
        )

        if result['generated'] == 0:
            return Response({
                'generated': 0,
                'total_units_committed': 0,
                'message': 'No shortage items found. All requirements are adequately covered.',
            }, status=status.HTTP_200_OK)

        return Response({
            'generated': result['generated'],
            'total_units_committed': result['total_units_committed'],
            'message': (
                f"Auto-scheduled {result['generated']} delivery commitments "
                f"covering all deficits "
                f"({result['total_units_committed']:,} total units)."
            ),
        }, status=status.HTTP_200_OK)


# ─────────────────────────────────────────────────────────────────────────────
# View 6: Change Log / History Drawer
# GET /api/vendor-delivery-schedules/change-logs/
# ─────────────────────────────────────────────────────────────────────────────

class DeliveryChangeLogListView(APIView):
    permission_classes = [AllowAny]

    def get(self, request):
        po_number = request.query_params.get('po_number')
        comp_code = request.query_params.get('component_code') or request.query_params.get('componentCode')
        vendor_name = request.query_params.get('vendor_name') or request.query_params.get('vendorName')
        changed_by = request.query_params.get('changed_by') or request.query_params.get('changedBy')
        schedule_id = request.query_params.get('schedule_id') or request.query_params.get('scheduleId')
        date_from = request.query_params.get('date_from')
        date_to = request.query_params.get('date_to')
        field_type = request.query_params.get('field_type')  # DATE|QTY|STATUS
        search = request.query_params.get('search', '').strip()
        try:
            page = max(1, int(request.query_params.get('page', 1)))
        except (ValueError, TypeError):
            page = 1
        try:
            page_size = max(1, min(200, int(request.query_params.get('page_size', 50))))
        except (ValueError, TypeError):
            page_size = 50

        qs = DeliveryScheduleChangeLog.objects.all().order_by('-changed_at')

        if po_number:
            qs = qs.filter(po_number__iexact=po_number)
        if comp_code:
            qs = qs.filter(component_code__iexact=comp_code)
        if vendor_name:
            qs = qs.filter(vendor_name__icontains=vendor_name)
        if changed_by:
            qs = qs.filter(changed_by__icontains=changed_by)
        if schedule_id:
            qs = qs.filter(schedule_id=schedule_id)
        if date_from:
            qs = qs.filter(changed_at__date__gte=date_from)
        if date_to:
            qs = qs.filter(changed_at__date__lte=date_to)

        # field_type filter (maps to HistoryDrawer category filter)
        if field_type == 'DATE':
            qs = qs.filter(
                Q(field_changed__icontains='date') |
                Q(field_changed__icontains='arrival')
            )
        elif field_type == 'QTY':
            qs = qs.filter(
                Q(field_changed__icontains='qty') |
                Q(field_changed__icontains='quantity')
            )
        elif field_type == 'STATUS':
            qs = qs.filter(
                Q(field_changed__icontains='status') |
                Q(field_changed__icontains='cancelled')
            )

        if search:
            qs = qs.filter(
                Q(po_number__icontains=search) |
                Q(component_code__icontains=search) |
                Q(component_description__icontains=search) |
                Q(vendor_name__icontains=search) |
                Q(changed_by__icontains=search) |
                Q(field_changed__icontains=search) |
                Q(reason_for_change__icontains=search) |
                Q(old_value__icontains=search) |
                Q(new_value__icontains=search)
            )

        total = qs.count()
        offset = (page - 1) * page_size
        page_qs = qs[offset: offset + page_size]

        serializer = DeliveryScheduleChangeLogSerializer(page_qs, many=True)
        return Response({
            'count': total,
            'page': page,
            'page_size': page_size,
            'total_pages': (total + page_size - 1) // page_size if total > 0 else 1,
            'results': serializer.data,
        }, status=status.HTTP_200_OK)


# ─────────────────────────────────────────────────────────────────────────────
# View 7: Download Blank Excel Template
# GET /api/exports/vendor-schedule-blank-template/?month=YYYY-MM
# ─────────────────────────────────────────────────────────────────────────────

class VendorScheduleBlankTemplateView(APIView):
    permission_classes = [AllowAny]

    def get(self, request):
        import openpyxl
        month = request.query_params.get('month', '')
        week_no = request.query_params.get('week_no', '1')
        week_start = request.query_params.get('week_start', f'{month}-10' if month else '2026-08-10')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Schedule_Template'

        headers = [
            'PO_Number', 'Component_Code', 'Component_Description',
            'Vendor_Code', 'Vendor_Name', 'Buyer_Name',
            'Expected_Delivery_Date', 'Week_No', 'Promised_Quantity',
            'Delivery_Status', 'Carrier_Or_Tracking', 'Notes'
        ]
        ws.append(headers)

        clean_month = month.replace("-", "") if month else "202608"

        # Two sample rows
        ws.append([
            f'PO-{clean_month}-8801', '100201',
            'Die-Cast Aluminum Housing (Panther)',
            'V-1020', 'CastAlu Technologies GmbH', 'Rajesh Kumar',
            week_start, int(week_no) if str(week_no).isdigit() else 1, 3500,
            'CONFIRMED_ON_TRACK', 'Direct Express Truck (TN-04-AB-9821)',
            'Confirmed by vendor for Monday morning dock receipt'
        ])
        ws.append([
            f'PO-{clean_month}-8802', '200405',
            'Composite Carbon Sliding Vane (3-Set)',
            'V-1030', 'CarbonTech Materials AG', 'Ananya Sharma',
            week_start, int(week_no) if str(week_no).isdigit() else 1, 4000,
            'CONFIRMED_ON_TRACK', 'DHL Freight Logistics', 'Weekly batch shipment'
        ])

        column_widths = [16, 16, 38, 14, 28, 18, 24, 10, 18, 22, 30, 40]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        filename = f'SAP_Vendor_Delivery_Schedule_Blank_Template_{month or "template"}.xlsx'
        response = StreamingHttpResponse(
            buffer,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


# ─────────────────────────────────────────────────────────────────────────────
# View 8: Download Pre-Filled Template with Computed Requirements
# GET /api/exports/vendor-schedule-prefilled/?month=&scope_mode=&week_code=
# ─────────────────────────────────────────────────────────────────────────────

class VendorSchedulePrefilledTemplateView(APIView):
    permission_classes = [AllowAny]

    def get(self, request):
        import openpyxl
        month = request.query_params.get('month')
        scope_mode = request.query_params.get('scope_mode', 'WEEK')
        week_code = request.query_params.get('week_code')

        if not month:
            return Response({'error': 'month is required'}, status=status.HTTP_400_BAD_REQUEST)

        matrix = ConsolidatedMatrixService.compute(
            month=month,
            scope_mode=scope_mode,
            selected_week_code=week_code,
        )

        selected_week = matrix.get('selectedWeek')
        week_start = selected_week['startDate'] if selected_week else f'{month}-10'
        week_no_val = selected_week['weekNo'] if selected_week else 1
        scope_label = f"Week_{week_no_val}" if scope_mode == 'WEEK' else 'Full_Month'

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Consolidated_Requirements'

        headers = [
            'PO_Number', 'Component_Code', 'Component_Description', 'Category',
            'Is_Common_Part', 'Vendor_Code', 'Vendor_Name', 'Buyer_Name',
            'Lead_Time_Days', 'MB52_Available_Stock', 'Gross_Requirement',
            'Current_Scheduled_Inward', 'Deficit_Shortage',
            'Expected_Delivery_Date', 'Week_No', 'Promised_Quantity',
            'Delivery_Status', 'Carrier_Or_Tracking', 'Notes'
        ]
        ws.append(headers)

        clean_month = month.replace("-", "")
        for idx, row in enumerate(matrix['rows']):
            suggested_qty = (
                row['selectedScopeDeficit']
                if row['hasShortage'] else
                max(row['selectedScopeGrossReq'], 1000)
            )
            common_label = (
                f"YES ({row['sharedInFGsCount']} FGs)"
                if row['isCommonPart'] else 'NO'
            )
            ws.append([
                f'PO-{clean_month}-{9000 + idx}',
                row['componentCode'],
                row['componentDescription'],
                row['category'],
                common_label,
                row['vendorCode'],
                row['vendorName'],
                row['buyerName'],
                row['leadTimeDays'],
                row['availableStock'],
                row['selectedScopeGrossReq'],
                row['selectedScopeInward'],
                row['selectedScopeDeficit'],
                week_start,
                week_no_val,
                suggested_qty,
                'CONFIRMED_ON_TRACK',
                'Road Logistics Standard',
                (
                    f"Shared in {row['sharedInFGsCount']} FGs — maintain priority inward."
                    if row['isCommonPart']
                    else 'Standard production replenishment'
                ),
            ])

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        filename = f'SAP_RM_Requirements_PreFilled_{month}_{scope_label}.xlsx'
        response = StreamingHttpResponse(
            buffer,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


# ─────────────────────────────────────────────────────────────────────────────
# View 9: Export Consolidated Matrix to Excel Report
# GET /api/exports/vendor-schedule-matrix/?month=&buyer=&category=&criticality=
# ─────────────────────────────────────────────────────────────────────────────

class VendorScheduleMatrixExportView(APIView):
    permission_classes = [AllowAny]

    def get(self, request):
        import openpyxl
        from openpyxl.styles import PatternFill

        month = request.query_params.get('month')
        scope_mode = request.query_params.get('scope_mode', 'MONTH')
        week_code = request.query_params.get('week_code')
        buyer = request.query_params.get('buyer', 'ALL')
        category = request.query_params.get('category', 'ALL')
        criticality = request.query_params.get('criticality')

        if not month:
            return Response({'error': 'month is required'}, status=status.HTTP_400_BAD_REQUEST)

        matrix = ConsolidatedMatrixService.compute(
            month=month,
            scope_mode=scope_mode,
            selected_week_code=week_code,
            buyer_filter=buyer,
            category_filter=category,
            criticality_filter=criticality,
        )
        month_weeks = matrix['monthWeeks']

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Consolidated_RM_Matrix'

        fixed_headers = [
            'Component Code', 'Description', 'Category', 'UOM',
            'Is Common Item', 'Buyer', 'Vendor', 'Vendor Code',
            'MB52 Physical Stock', 'Frozen Reserved', 'Available Stock'
        ]
        week_headers = []
        for w in month_weeks:
            week_headers += [
                f"{w['weekLabel']} Gross Req",
                f"{w['weekLabel']} Inward Inbound",
            ]
        tail_headers = [
            'Total Month Gross Demand', 'Total Month Inward Scheduled',
            'Ending Balance', 'Current Status'
        ]
        ws.append(fixed_headers + week_headers + tail_headers)

        red_fill = PatternFill('solid', fgColor='FFCDD2')
        green_fill = PatternFill('solid', fgColor='C8E6C9')

        for row_data in matrix['rows']:
            row = [
                row_data['componentCode'],
                row_data['componentDescription'],
                row_data['category'],
                row_data['uom'],
                f"Shared in {row_data['sharedInFGsCount']} FGs" if row_data['isCommonPart'] else 'Single FG',
                row_data['buyerName'],
                row_data['vendorName'],
                row_data['vendorCode'],
                row_data['totalPhysicalStock'],
                row_data['reservedStock'],
                row_data['availableStock'],
            ]
            for w in month_weeks:
                row.append(row_data['weeklyGrossReq'].get(w['weekCode'], 0))
                row.append(row_data['weeklyInwardDeliveries'].get(w['weekCode'], 0))
            row += [
                row_data['monthTotalGrossReq'],
                row_data['monthTotalInward'],
                row_data['monthEndingBalance'],
                'SHORTAGE_DEFICIT' if row_data['hasShortage'] else 'ADEQUATE',
            ]
            ws.append(row)
            status_cell = ws.cell(
                row=ws.max_row, column=len(fixed_headers + week_headers + tail_headers)
            )
            status_cell.fill = red_fill if row_data['hasShortage'] else green_fill

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        filename = f'Consolidated_RM_Schedule_Matrix_{month}.xlsx'
        response = StreamingHttpResponse(
            buffer,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
